"""HTTP surface for the options model.

    GET  /api/options-model/model        full payload (ticker via ?ticker=)
    GET  /api/options-model/quick        payload without chain or backtest
    GET  /api/options-model/backtest     walk-forward validation only
    POST /api/options-model/refresh      drop the cache for one ticker
    GET  /api/options-model/export       Excel workbook of the current view

Query parameters on ``/model``:

    ticker     symbol (default TSEM)
    horizon    simulation horizon in trading days (default 21)
    engine     gbm | student_t | garch_t | bootstrap
    drift      risk_neutral | zero | historical | custom
    mu         annual simple return, only used when drift=custom
    paths      Monte Carlo paths (2,000 - 100,000)
"""

import io
import logging
from datetime import date

from flask import Blueprint, jsonify, request, send_file

from . import data as market
from . import service

logger = logging.getLogger(__name__)

options_model_bp = Blueprint('options_model', __name__)


def _params():
    args = request.args
    try:
        custom = args.get('mu')
        custom = float(custom) if custom not in (None, '') else None
    except ValueError:
        custom = None
    try:
        paths = int(args.get('paths', service.DEFAULT_PATHS))
    except ValueError:
        paths = service.DEFAULT_PATHS
    try:
        horizon = int(args.get('horizon', service.DEFAULT_HORIZON))
    except ValueError:
        horizon = service.DEFAULT_HORIZON
    return {
        'ticker': (args.get('ticker') or market.DEFAULT_TICKER).upper().strip()[:12],
        'horizon': horizon,
        'engine': args.get('engine', service.DEFAULT_ENGINE),
        'drift_mode': args.get('drift', service.DEFAULT_DRIFT),
        'n_paths': paths,
        'custom_drift': custom,
    }


@options_model_bp.route('/model')
def model():
    p = _params()
    payload = service.get_model(**p, with_backtest=True)
    if payload is None:
        return jsonify({'error': f"No price history for {p['ticker']}."}), 404
    return jsonify(payload)


@options_model_bp.route('/quick')
def quick():
    """Volatility and simulation only — no chain fetch, no backtest.

    Used when flipping the engine or drift control, where re-fetching the
    chain and re-running the walk-forward would add seconds for numbers that
    did not change.
    """
    p = _params()
    payload = service.get_model(**p, with_backtest=False)
    if payload is None:
        return jsonify({'error': f"No price history for {p['ticker']}."}), 404
    return jsonify(payload)


@options_model_bp.route('/backtest')
def backtest():
    p = _params()
    payload = service.get_model(**p, with_backtest=True)
    if payload is None:
        return jsonify({'error': f"No price history for {p['ticker']}."}), 404
    return jsonify(payload.get('backtest') or {'available': False})


@options_model_bp.route('/refresh', methods=['POST'])
def refresh():
    ticker = (request.args.get('ticker') or '').upper().strip() or None
    service.clear_cache(ticker)
    return jsonify({'status': 'ok', 'cleared': ticker or 'all'})


@options_model_bp.route('/export')
def export():
    p = _params()
    payload = service.get_model(**p, with_backtest=True)
    if payload is None:
        return jsonify({'error': f"No price history for {p['ticker']}."}), 404
    return _export_excel(payload)


def _export_excel(payload):
    """Workbook: summary, vol, direction, chain, structures, calibration."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF', size=10)
    head_fill = PatternFill('solid', fgColor='1F2937')
    title_font = Font(bold=True, size=13)
    pct = '0.0%'
    money = '#,##0.00'

    def sheet(ws, title, headers, rows, widths=None, formats=None):
        ws['A1'] = title
        ws['A1'].font = title_font
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=3, column=i, value=h)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        for ri, row in enumerate(rows, start=4):
            for ci, val in enumerate(row, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                if formats and formats.get(ci):
                    cell.number_format = formats[ci]
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = 'A4'

    meta = payload['meta']
    sim = payload['simulation']
    primary = sim['primary']

    ws = wb.active
    ws.title = 'Summary'
    rows = [
        ['Ticker', meta['ticker']],
        ['Name', meta['name']],
        ['Spot', meta['spot']],
        ['As of', meta['as_of']],
        ['Risk-free (cc)', meta['risk_free']],
        ['Next earnings', meta.get('earnings_date') or 'n/a'],
        ['Engine', sim['settings']['engine_label']],
        ['Horizon (trading days)', sim['settings']['horizon_days']],
        ['Drift convention', sim['settings']['drift_mode']],
        ['Paths', sim['settings']['n_paths']],
        ['Model vol (horizon)', sim['settings']['sigma_ann']],
        ['P(up)', primary['p_up']],
        ['Median return', primary['median_return']],
        ['Expected abs move', primary['expected_abs_move']],
        ['68% band low', primary['band68'][0]],
        ['68% band high', primary['band68'][1]],
        ['CVaR 5%', primary['cvar05']],
        [],
        ['Narrative', payload.get('narrative', '')],
    ]
    sheet(ws, f"{meta['ticker']} — options model", ['Field', 'Value'], rows,
          widths=[28, 90])
    ws['B19'].alignment = Alignment(wrap_text=True, vertical='top')

    ws = wb.create_sheet('Volatility')
    v = payload['vol']
    rv = v['realized']
    vrows = [[k, val] for k, val in rv.items() if val is not None]
    if v.get('garch'):
        g = v['garch']
        vrows += [[], ['GARCH alpha', g['alpha']], ['GARCH beta', g['beta']],
                  ['Persistence', g['persistence']],
                  ['Half-life (days)', g['half_life_days']],
                  ['Long-run vol', g['long_run_vol']],
                  ['Spot vol', g['spot_vol']], ['Student-t df', g['nu']]]
    sheet(ws, 'Volatility estimates (annualized)', ['Estimator', 'Value'],
          vrows, widths=[28, 16], formats={2: '0.000'})

    ws = wb.create_sheet('Direction')
    sheet(ws, 'Direction probability by horizon',
          ['Days', 'Calendar days', 'Vol', 'P(up)', 'P(down)', 'Median ret',
           '1sd move', '68% low', '68% high', 'Variance drag'],
          [[d['days'], d['calendar_days'], d['vol'], d['p_up'], d['p_down'],
            d['median_return'], d['one_sd_move'], d['band68'][0], d['band68'][1],
            d['variance_drag']] for d in payload['direction']],
          widths=[8, 14, 10, 10, 10, 12, 10, 10, 10, 13],
          formats={3: '0.0%', 4: pct, 5: pct, 6: pct, 7: pct, 8: money,
                   9: money, 10: pct})

    ws = wb.create_sheet('Chain vol')
    sheet(ws, 'Implied vs model volatility by expiry',
          ['Expiry', 'DTE', 'Trading days', 'ATM IV', 'Model vol', 'VRP',
           'IV/Model', '25d put IV', '25d call IV', 'Skew', 'Implied move %',
           'Model move %'],
          [[c['expiry'], c['dte'], c['trading_days'], c['atm_iv'], c['model_vol'],
            c['vrp'], c['vrp_ratio'], c['iv_25d_put'], c['iv_25d_call'],
            c['skew_25d'], c['implied_move_pct'], c['model_move_pct']]
           for c in payload.get('chain_vol') or []],
          widths=[12, 7, 13, 10, 11, 9, 10, 11, 12, 9, 14, 13],
          formats={4: pct, 5: pct, 6: pct, 7: '0.00', 8: pct, 9: pct, 10: pct,
                   11: '0.0', 12: '0.0'})

    ws = wb.create_sheet('Structures')
    srows = []
    for s in (payload.get('strategies') or {}).get('all') or []:
        stress = s.get('ev_stress') or {}
        srows.append([
            s['name'], s['expiry'], s['dte'], s['type'],
            ', '.join(f"{l['qty']:+d} {l['kind']} {l['strike']:g}"
                      for l in s['legs'] if l['kind'] != 'stock'),
            s['net_debit'], s['ev'], s['ev_marketable'],
            stress.get('vol_down'), stress.get('vol_up'),
            s['pop'], s['capital'], s['max_loss'], s['max_profit'],
            s['ev_on_capital'], s['kelly'], s['greeks']['delta'],
            s['greeks']['vega'], s['greeks']['theta'], s['avg_iv'],
            s['liquidity'], 'yes' if s.get('robust') else 'no',
        ])
    sheet(ws, 'Option structures scored against the simulated distribution',
          ['Structure', 'Expiry', 'DTE', 'Type', 'Legs', 'Net debit', 'EV',
           'EV at market fill', 'EV if vol low', 'EV if vol high', 'POP',
           'Capital', 'Max loss', 'Max profit', 'EV/capital', 'Kelly',
           'Delta', 'Vega', 'Theta', 'Avg IV', 'Liquidity', 'Robust'],
          srows,
          widths=[32, 12, 6, 18, 30, 11, 10, 15, 13, 14, 8, 10, 10, 11, 11, 8,
                  9, 9, 9, 9, 11, 9],
          formats={6: money, 7: money, 8: money, 9: money, 10: money, 11: pct,
                   12: money, 13: money, 14: money, 15: '0.00%', 16: '0.00',
                   17: '0.000', 18: '0.000', 19: '0.000', 20: pct})

    ws = wb.create_sheet('Strike edges')
    sheet(ws, 'Model vs market probability of finishing in the money',
          ['Kind', 'Expiry DTE', 'Strike', 'Moneyness', 'IV', 'Delta', 'Mid',
           'P(ITM) market', 'P(ITM) model', 'Edge', 'Model price', 'Price edge',
           'Open interest'],
          [[e['kind'], e['dte'], e['strike'], e['moneyness'], e['iv'],
            e['delta'], e['mid'], e['p_itm_market'], e['p_itm_model'],
            e['prob_edge'], e['model_price'], e['price_edge'],
            e['open_interest']] for e in payload.get('strike_edges') or []],
          widths=[7, 11, 9, 11, 9, 9, 9, 14, 13, 9, 12, 11, 13],
          formats={4: '0.000', 5: pct, 6: '0.000', 7: money, 8: pct, 9: pct,
                   10: pct, 11: money, 12: money})

    bt = payload.get('backtest') or {}
    if bt.get('available'):
        ws = wb.create_sheet('Calibration')
        d = bt.get('direction') or {}
        rows = [
            ['Windows tested', bt.get('n_windows')],
            ['Horizon (days)', bt.get('horizon_days')],
            ['Brier — model', d.get('brier_model')],
            ['Brier — coin flip', d.get('brier_coinflip')],
            ['Brier — momentum tilt', d.get('brier_momentum')],
            ['Directional skill vs coin flip', d.get('skill_vs_coinflip')],
            [],
            ['Vol model', 'RMSE', 'MAE', 'Bias', 'R2', 'MZ slope'],
        ]
        for name, stats in (bt.get('volatility') or {}).items():
            rows.append([name, stats['rmse'], stats['mae'], stats['bias'],
                         stats['r2'], stats['mz_slope']])
        rows.append([])
        rows.append(['Interval', 'Nominal', 'Empirical'])
        for iv in (bt.get('coverage') or {}).get('intervals') or []:
            rows.append(['coverage', iv['nominal'], iv['empirical']])
        rows += [[], ['Verdict', bt.get('verdict', '')]]
        sheet(ws, 'Walk-forward validation', ['Metric', 'Value', '', '', '', ''],
              rows, widths=[32, 14, 12, 12, 10, 12])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{meta['ticker'].lower()}_options_model_{date.today()}.xlsx")
